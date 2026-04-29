from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from bittensor_collector.models import EvidenceItem, SubnetRecord


def build_field_dictionary() -> list[dict]:
    return [
        {"字段": "注册费(TAO)", "说明": "子网注册成本，单位 TAO", "规则": "原值输出"},
        {"字段": "中游日收益(TAO)", "说明": "矿工收益中位数或P30", "规则": "由income_mode决定"},
        {"字段": "Gini", "说明": "收益或份额分布不平等程度", "规则": "若缺失则从分布计算"},
        {"字段": "HHI", "说明": "份额集中度指数", "规则": "若缺失则从分布计算"},
        {"字段": "热度", "说明": "由矿工数与总质押量组合并按全表归一化后的100分制评分", "规则": "最高值为100"},
        {"字段": "初步结论", "说明": "阈值规则生成的风险提示", "规则": "见配置thresholds"},
    ]


def export_excel(records: list[SubnetRecord], evidence_items: list[EvidenceItem], output_path: str) -> None:
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    main_rows = [r.to_main_row() for r in records]
    evidence_rows = [
        {
            "netuid": item.netuid,
            "source": item.source,
            "field_name": item.field_name,
            "source_url": item.source_url,
            "raw_text": item.raw_text,
            "manual_review_needed": item.manual_review_needed,
        }
        for item in evidence_items
    ]

    workbook = Workbook()
    main_sheet = workbook.active
    main_sheet.title = "主表"
    _write_sheet(main_sheet, main_rows)

    evidence_sheet = workbook.create_sheet("原始证据")
    _write_sheet(evidence_sheet, evidence_rows)

    dictionary_sheet = workbook.create_sheet("字段说明")
    _write_sheet(dictionary_sheet, build_field_dictionary())

    workbook.save(output_path)


def _write_sheet(worksheet, rows: list[dict]) -> None:
    if not rows:
        worksheet.append(["empty"])
        return

    headers = list(rows[0].keys())
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])
