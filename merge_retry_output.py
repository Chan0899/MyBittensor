from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="将重跑结果按子网ID合并回原始 Excel")
    parser.add_argument("--base", required=True, help="原始 Excel 路径")
    parser.add_argument("--retry", required=True, help="重跑结果 Excel 路径")
    return parser.parse_args()


def _sheet_to_rows(workbook_path: Path, sheet_name: str) -> tuple[list[str], list[dict]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in {workbook_path}")
        worksheet = workbook[sheet_name]
        row_iter = worksheet.iter_rows(values_only=True)
        headers_row = next(row_iter, None)
        if not headers_row:
            return [], []
        headers = [str(value).strip() if value is not None else "" for value in headers_row]
        rows: list[dict] = []
        for values in row_iter:
            if values is None:
                continue
            row = {
                headers[index]: values[index] if index < len(values) else None
                for index in range(len(headers))
                if headers[index]
            }
            rows.append(row)
        return headers, rows
    finally:
        workbook.close()


def _replace_rows_by_key(base_rows: list[dict], retry_rows: list[dict], key: str) -> list[dict]:
    retry_map = {row[key]: row for row in retry_rows if row.get(key) is not None}
    merged_rows: list[dict] = []
    replaced_keys: set[object] = set()

    for row in base_rows:
        row_key = row.get(key)
        if row_key in retry_map:
            merged_rows.append(retry_map[row_key])
            replaced_keys.add(row_key)
        else:
            merged_rows.append(row)

    for row in retry_rows:
        row_key = row.get(key)
        if row_key is None or row_key in replaced_keys:
            continue
        merged_rows.append(row)

    return merged_rows


def _write_sheet(worksheet, headers: list[str], rows: list[dict]) -> None:
    worksheet.delete_rows(1, worksheet.max_row)
    if not headers:
        worksheet.append(["empty"])
        return
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])


def main() -> int:
    args = parse_args()
    base_path = Path(args.base)
    retry_path = Path(args.retry)

    if not base_path.exists():
        print(f"ERROR: Base Excel file not found: {base_path}")
        return 1
    if not retry_path.exists():
        print(f"ERROR: Retry Excel file not found: {retry_path}")
        return 1

    main_headers, base_main_rows = _sheet_to_rows(base_path, "主表")
    _, retry_main_rows = _sheet_to_rows(retry_path, "主表")
    _, base_evidence_rows = _sheet_to_rows(base_path, "原始证据")
    _, retry_evidence_rows = _sheet_to_rows(retry_path, "原始证据")

    if "子网ID" not in main_headers:
        print("ERROR: Column '子网ID' not found in 主表")
        return 1

    merged_main_rows = _replace_rows_by_key(base_main_rows, retry_main_rows, "子网ID")
    retry_netuids = {row.get("子网ID") for row in retry_main_rows if row.get("子网ID") is not None}
    merged_evidence_rows = [
        row for row in base_evidence_rows if row.get("netuid") not in retry_netuids
    ]
    merged_evidence_rows.extend(retry_evidence_rows)

    base_workbook = load_workbook(base_path)
    try:
        _write_sheet(base_workbook["主表"], main_headers, merged_main_rows)
        evidence_headers = [cell.value for cell in base_workbook["原始证据"][1]] if base_workbook["原始证据"].max_row >= 1 else []
        _write_sheet(base_workbook["原始证据"], evidence_headers, merged_evidence_rows)
        base_workbook.save(base_path)
    finally:
        base_workbook.close()

    print(f"Merged retry result into {base_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())