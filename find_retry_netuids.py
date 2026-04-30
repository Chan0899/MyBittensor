from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def _is_manual_review(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    text = str(value).strip().lower()
    return text in {"1", "true", "yes", "y", "是"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="从导出的 Excel 中提取需重跑的 netuid 列表")
    parser.add_argument("--xlsx", required=True, help="Excel 文件路径")
    parser.add_argument("--count", action="store_true", help="仅输出需重跑的子网个数")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook_path = Path(args.xlsx)
    if not workbook_path.exists():
        print(f"ERROR: Excel file not found: {workbook_path}")
        return 1

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "主表" not in workbook.sheetnames:
            print("ERROR: Sheet '主表' not found")
            return 1

        worksheet = workbook["主表"]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            print("ERROR: Sheet '主表' is empty")
            return 1

        header_map = {str(name).strip(): index for index, name in enumerate(headers) if name is not None}
        if "子网ID" not in header_map or "需人工复核" not in header_map:
            print("ERROR: Required columns '子网ID' or '需人工复核' are missing")
            return 1

        netuid_index = header_map["子网ID"]
        manual_review_index = header_map["需人工复核"]
        retry_netuids: list[str] = []

        for row in rows:
            if row is None:
                continue
            netuid = row[netuid_index] if netuid_index < len(row) else None
            manual_review_needed = row[manual_review_index] if manual_review_index < len(row) else None
            if netuid is None or not _is_manual_review(manual_review_needed):
                continue
            retry_netuids.append(str(int(netuid)))

        if args.count:
            print(len(retry_netuids))
        else:
            print(",".join(retry_netuids))
        return 0
    finally:
        workbook.close()


if __name__ == "__main__":
    raise SystemExit(main())